"""In-memory text extraction for CV files."""

from io import BytesIO
from pathlib import PurePath
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader


MAX_CV_TEXT_CHARS = 60_000


class CVDocumentError(ValueError):
    """A CV file is unsupported, invalid, or can't be read as text."""


def extract_cv_text(filename: str, content: bytes) -> str:
    """Extract text from a PDF or DOCX file without saving it to disk."""

    extension = PurePath(filename).suffix.lower()
    if extension == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise CVDocumentError("That file doesn't appear to be a valid PDF.")
        try:
            reader = PdfReader(BytesIO(content), strict=False)
            if reader.is_encrypted:
                raise CVDocumentError("Password-protected PDFs aren't supported.")
            if len(reader.pages) > 20:
                raise CVDocumentError("Please upload a PDF with no more than 20 pages.")
            text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        except CVDocumentError:
            raise
        except Exception as error:
            raise CVDocumentError("Couldn't read that PDF. Please check that it is valid.") from error

        if not text.strip():
            raise CVDocumentError(
                "This PDF has no selectable text. Scanned PDFs need OCR and aren't supported yet."
            )
    elif extension == ".docx":
        try:
            with ZipFile(BytesIO(content)) as archive:
                entries = archive.infolist()
                if "word/document.xml" not in {entry.filename for entry in entries}:
                    raise CVDocumentError("That file doesn't appear to be a valid DOCX document.")
                if sum(entry.file_size for entry in entries) > 25 * 1024 * 1024:
                    raise CVDocumentError("The DOCX file expands to too much data to process.")
            document = Document(BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        except CVDocumentError:
            raise
        except Exception as error:
            raise CVDocumentError("Couldn't read that DOCX. Please check that it is valid.") from error
    else:
        raise CVDocumentError("Please attach a PDF or DOCX file.")

    if not text.strip():
        raise CVDocumentError("No text was found in that CV file.")
    if len(text) > MAX_CV_TEXT_CHARS:
        raise CVDocumentError("The extracted CV is too long. Please upload a shorter document.")
    return text
