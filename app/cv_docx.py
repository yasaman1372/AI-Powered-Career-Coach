import re
from io import BytesIO

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from app.cv_pdf import BULLET, _is_heading

ACCENT = RGBColor(0x1F, 0x5B, 0x4A)
MUTED = RGBColor(0x5B, 0x66, 0x60)
LINK = re.compile(
    r"(https?://[^\s,;|]+|www\.[^\s,;|]+|[\w.+-]+@[\w-]+\.[\w.-]+"
    r"|(?:linkedin|github)\.com/[^\s,;|]+)",
    re.IGNORECASE,
)
BOLD = re.compile(r"\*\*(.+?)\*\*")


def _add_hyperlink(paragraph, text: str, size: Pt | None = None) -> None:
    if "@" in text and not text.lower().startswith("http"):
        target = "mailto:" + text
    elif text.lower().startswith(("http://", "https://")):
        target = text
    else:
        target = "https://" + text
    rel_id = paragraph.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5B4A")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([style, color, underline])
    if size:
        half_points = OxmlElement("w:sz")
        half_points.set(qn("w:val"), str(int(size.pt * 2)))
        props.append(half_points)
    run.append(props)
    content = OxmlElement("w:t")
    content.text = text
    content.set(qn("xml:space"), "preserve")
    run.append(content)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_text(paragraph, text: str, size: Pt | None = None, color: RGBColor | None = None) -> None:
    """Add text with **bold** spans and clickable links."""

    for bold_index, bold_part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        for piece in LINK.split(bold_part):
            if not piece:
                continue
            if LINK.fullmatch(piece):
                _add_hyperlink(paragraph, piece.rstrip(".)"), size)
                continue
            run = paragraph.add_run(piece)
            run.bold = bold_index % 2 == 1
            if size:
                run.font.size = size
            if color:
                run.font.color.rgb = color


def _bottom_border(paragraph) -> None:
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "1F5B4A")):
        bottom.set(qn(f"w:{key}"), value)
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def build_cv_docx(text: str) -> bytes:
    """Render plain-text CV content as an editable Word document."""

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(20)
    section.top_margin = section.bottom_margin = Mm(18)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(3)

    lines = [line.rstrip() for line in text.replace("\r\n", "\n").split("\n")]
    while lines and not lines[0].strip():
        lines.pop(0)

    if lines:
        name = document.add_paragraph()
        run = name.add_run(BOLD.sub(r"\1", lines.pop(0).lstrip("# ").strip()))
        run.bold = True
        run.font.size = Pt(22)
        run.font.color.rgb = ACCENT
    contact = []
    while lines and lines[0].strip() and not _is_heading(lines[0]) and not BULLET.match(lines[0]):
        contact.append(lines.pop(0).strip())
    if contact:
        _add_text(document.add_paragraph(), " | ".join(contact), Pt(9.5), MUTED)

    for line in lines:
        if not line.strip():
            continue
        match = BULLET.match(line)
        if match:
            _add_text(document.add_paragraph(style="List Bullet"), match.group(1).strip())
        elif _is_heading(line):
            heading = document.add_paragraph()
            heading.paragraph_format.space_before = Pt(10)
            run = heading.add_run(line.strip().lstrip("#").strip().rstrip(":"))
            run.bold = True
            run.font.size = Pt(11.5)
            run.font.color.rgb = ACCENT
            _bottom_border(heading)
        else:
            _add_text(document.add_paragraph(), line.strip())

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
